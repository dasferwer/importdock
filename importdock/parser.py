import csv
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook


def rows(path, format):
    if format == "csv":
        with open(path, encoding="utf-8-sig", newline="") as stream:
            yield from csv.reader(stream)
    else:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            for row in workbook.active.iter_rows(values_only=True):
                yield ["" if value is None else str(value) for value in row]
        finally:
            workbook.close()


def columns(header, mapping):
    if len(set(header)) != len(header):
        raise ValueError("Заголовки колонок должны быть уникальными")
    if set(mapping) != {"external_id", "amount"}:
        raise ValueError("Нужно сопоставить external_id и amount")
    if mapping["external_id"] == mapping["amount"]:
        raise ValueError("Поля должны использовать разные колонки")
    try:
        return {key: header.index(value) for key, value in mapping.items()}
    except ValueError as exc:
        raise ValueError("В файле нет указанной колонки") from exc


def validate(row, indices):
    try:
        identity = row[indices["external_id"]].strip()
        raw = row[indices["amount"]].strip()
        amount = Decimal(raw)
    except (IndexError, InvalidOperation) as exc:
        raise ValueError("Нет поля или сумма не является числом") from exc
    if not identity or len(identity) > 120 or "\x00" in identity:
        raise ValueError("Идентификатор должен содержать от 1 до 120 символов")
    if (
        not amount.is_finite()
        or abs(amount) >= Decimal("1e16")
        or amount != amount.quantize(Decimal("0.01"))
    ):
        raise ValueError("Нужна конечная сумма с точностью не более двух знаков")
    return identity, amount


class PhysicalLines:
    def __init__(self, stream):
        self.stream = stream
        self.record_bytes = 0

    def __iter__(self):
        return self

    def __next__(self):
        raw = self.stream.readline(1024 * 1024 + 1)
        if not raw:
            raise StopIteration
        self.record_bytes += len(raw)
        if len(raw) > 1024 * 1024 or self.record_bytes > 2 * 1024 * 1024:
            raise ValueError("Строка CSV превышает лимит размера")
        return raw.decode("utf-8-sig" if self.stream.tell() == len(raw) else "utf-8")


class SeekableCSV:
    """Позиция снимается после логической записи, включая поля с переводами строк."""

    def __init__(self, path, offset=0):
        self.stream = open(path, "rb")  # noqa: SIM115 — process закрывает поток в finally.
        self.stream.seek(offset)
        self.lines = PhysicalLines(self.stream)
        self.reader = csv.reader(self.lines, strict=True)

    def __iter__(self):
        return self

    def __next__(self):
        self.lines.record_bytes = 0
        try:
            return next(self.reader)
        except (csv.Error, UnicodeDecodeError) as exc:
            raise ValueError("Повреждён CSV: " + str(exc)) from exc

    @property
    def offset(self):
        return self.stream.tell()

    def seek(self, offset):
        self.stream.seek(offset)
        self.reader = csv.reader(self.lines, strict=True)

    def close(self):
        self.stream.close()
